# -*- coding: utf-8 -*-
"""
KML/KMZ -> Katman Bazında Shapefile

ArcPy ve GDAL kullanmaz. Yalnızca Python standart kütüphanesi ve pyshp gerekir.

Kurulum:
    python -m pip install pyshp

Kullanım:
    python kmz_to_shp_layers_separate.py "C:\\veri\\girdi.kml" "C:\\veri\\shp_cikti"

Davranış:
- Her KML/KMZ klasörü (katman) ayrı tutulur.
- Aynı geometri türündeki farklı katmanlar birleştirilmez.
- Dissolve, merge veya duplicate silme yapılmaz.
- Aynı koordinattaki iki obje dahi ayrı kayıtlarda kalır.
- MultiGeometry içindeki her geometri ayrı obje olarak yazılır.
- Nokta, çizgi ve poligonlar ayrı Shapefile olur.
- ExtendedData, SchemaData ve description içindeki HTML tablolar okunur.
- Türkçe öznitelik değerleri UTF-8 olarak korunur.
- Shapefile'ın 10 karakterlik alan adı sınırı nedeniyle okunabilir kısa teknik
  adlar üretilir; özgün adlar .qml alias ve *_field_map.csv içinde saklanır.
- Tam/kesilmemiş öznitelikler ayrıca *_attributes.csv ve *.jsonl dosyalarına yazılır.
- Doğrudan Excel (.xlsx) çıktısı üretilir: her kayıt için X, Y temsili koordinat
  ve tam geometri WKT (Well-Known Text) string olarak yazılır (POINT/LINESTRING/
  POLYGON). Bunun için openpyxl gerekir: python -m pip install openpyxl
  --no-excel bayrağı ile devre dışı bırakılabilir.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import unicodedata
import warnings
import zipfile
from collections import OrderedDict, defaultdict
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape as xml_escape

import shapefile

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    HAVE_OPENPYXL = True
except ImportError:  # openpyxl yoksa Excel çıktısı atlanır, uyarı verilir.
    HAVE_OPENPYXL = False


WGS84_PRJ = (
    'GEOGCS["WGS 84",DATUM["WGS_1984",'
    'SPHEROID["WGS 84",6378137,298.257223563]],'
    'PRIMEM["Greenwich",0],UNIT["degree",0.0174532925199433]]'
)

DBF_TEXT_LIMIT = 254
DBF_MAX_FIELDS = 250

GEOMETRY_SUFFIX = {
    "point": "nokta",
    "line": "cizgi",
    "polygon": "poligon",
}

TURKISH_CHAR_MAP = str.maketrans(
    {
        "Ç": "C", "Ğ": "G", "İ": "I", "Ö": "O", "Ş": "S", "Ü": "U",
        "ç": "c", "ğ": "g", "ı": "i", "ö": "o", "ş": "s", "ü": "u",
    }
)


# -----------------------------------------------------------------------------
# Temel yardımcılar
# -----------------------------------------------------------------------------

def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def clean_text(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return text.replace("\x00", "").strip()


def direct_child_text(element, wanted_name: str) -> str:
    for child in list(element):
        if local_name(child.tag) == wanted_name:
            return clean_text(child.text)
    return ""


def descendant_text(element, wanted_name: str) -> str:
    for child in element.iter():
        if local_name(child.tag) == wanted_name:
            return clean_text(child.text)
    return ""


def safe_file_part(value: str, fallback: str = "katman", max_length: int = 80) -> str:
    text = clean_text(value)
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", text)
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"_+", "_", text).strip(" ._")
    if not text:
        text = fallback
    return text[:max_length].rstrip(" ._") or fallback


def make_unique_output_name(seed: str, used_names: set[str]) -> str:
    candidate = safe_file_part(seed)
    key = candidate.casefold()
    if key not in used_names:
        used_names.add(key)
        return candidate

    counter = 2
    while True:
        suffix = "_{:02d}".format(counter)
        candidate2 = candidate[: 80 - len(suffix)] + suffix
        key = candidate2.casefold()
        if key not in used_names:
            used_names.add(key)
            return candidate2
        counter += 1


def readable_dbf_seed(original_name: str) -> str:
    """Özgün alan adından 10 karaktere uygun, okunabilir ASCII gövde üretir."""
    text = clean_text(original_name).translate(TURKISH_CHAR_MAP)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9_]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_").upper()
    if not text:
        text = "ALAN"
    if text[0].isdigit():
        text = "F_" + text
    return text


def unique_dbf_name(seed: str, used_names: set[str]) -> str:
    """DBF için en fazla 10 karakterlik benzersiz alan adı üretir."""
    base = readable_dbf_seed(seed)
    candidate = base[:10]
    if candidate.casefold() not in used_names:
        used_names.add(candidate.casefold())
        return candidate

    # Aynı ilk 10 karaktere sahip alanlar için kararlı kısa hash kullan.
    digest = hashlib.sha1(clean_text(seed).encode("utf-8")).hexdigest()[:3].upper()
    candidate = (base[:6] + "_" + digest)[:10]
    if candidate.casefold() not in used_names:
        used_names.add(candidate.casefold())
        return candidate

    counter = 2
    while True:
        suffix = str(counter)
        candidate = base[: 10 - len(suffix)] + suffix
        if candidate.casefold() not in used_names:
            used_names.add(candidate.casefold())
            return candidate
        counter += 1


# -----------------------------------------------------------------------------
# HTML tablo öznitelikleri
# -----------------------------------------------------------------------------

class HtmlTableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows = []
        self.current_row = None
        self.current_cell = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "tr":
            self.current_row = []
        elif tag in ("td", "th"):
            self.current_cell = []
        elif tag == "br" and self.current_cell is not None:
            self.current_cell.append("\n")

    def handle_data(self, data):
        if self.current_cell is not None:
            self.current_cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("td", "th") and self.current_cell is not None:
            value = clean_text(unescape("".join(self.current_cell))).replace("\xa0", " ")
            if self.current_row is not None:
                self.current_row.append(value)
            self.current_cell = None
        elif tag == "tr" and self.current_row is not None:
            if self.current_row:
                self.rows.append(self.current_row)
            self.current_row = None


def parse_description_table(html: str) -> OrderedDict:
    result = OrderedDict()
    if not html or "<" not in html:
        return result

    parser = HtmlTableParser()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return result

    for row in parser.rows:
        if len(row) < 2:
            continue
        key = clean_text(row[0])
        value = clean_text(" | ".join(part for part in row[1:] if part))
        if not key:
            continue
        if key in result and value and value != result[key]:
            result[key] = "{} | {}".format(result[key], value)
        elif key not in result:
            result[key] = value

    return result


# -----------------------------------------------------------------------------
# Geometri okuma
# -----------------------------------------------------------------------------

def parse_coordinates(text: str):
    coordinates = []
    for token in re.split(r"\s+", clean_text(text)):
        if not token:
            continue
        parts = token.split(",")
        if len(parts) < 2:
            continue
        try:
            lon = float(parts[0])
            lat = float(parts[1])
            alt = float(parts[2]) if len(parts) > 2 and parts[2] else 0.0
        except ValueError:
            continue
        coordinates.append([lon, lat, alt])
    return coordinates


def parse_gx_coord(text: str):
    parts = clean_text(text).split()
    if len(parts) < 2:
        return None
    try:
        lon = float(parts[0])
        lat = float(parts[1])
        alt = float(parts[2]) if len(parts) > 2 else 0.0
    except ValueError:
        return None
    return [lon, lat, alt]


def ring_area_xy(ring) -> float:
    if len(ring) < 3:
        return 0.0
    area = 0.0
    for index, point in enumerate(ring):
        next_point = ring[(index + 1) % len(ring)]
        area += point[0] * next_point[1] - next_point[0] * point[1]
    return area / 2.0


def close_ring(ring):
    if ring and ring[0][:2] != ring[-1][:2]:
        return ring + [ring[0][:]]
    return ring


def orient_polygon_rings(rings):
    """Shapefile standardına göre dış halka CW, iç halka CCW yapılır."""
    oriented = []
    for index, ring in enumerate(rings):
        ring = close_ring(ring)
        area = ring_area_xy(ring)
        if index == 0 and area > 0:
            ring = list(reversed(ring))
        elif index > 0 and area < 0:
            ring = list(reversed(ring))
        oriented.append(ring)
    return oriented


def geometry_properties(element) -> OrderedDict:
    properties = OrderedDict()
    for name in ("altitudeMode", "extrude", "tessellate"):
        value = direct_child_text(element, name)
        if value:
            properties["KML_{}".format(name)] = value
    return properties


def extract_geometries(element):
    """
    Dissolve veya merge yapmadan geometri örneklerini tek tek döndürür.

    Dönüş kayıtları:
        ("point", [x, y, z], properties)
        ("line", [[x, y, z], ...], properties)
        ("polygon", [ring1, ring2, ...], properties)
    """
    kind = local_name(element.tag)
    results = []

    if kind == "Point":
        properties = geometry_properties(element)
        for coordinate in parse_coordinates(descendant_text(element, "coordinates")):
            results.append(("point", coordinate, properties.copy()))

    elif kind == "LineString":
        coordinates = parse_coordinates(descendant_text(element, "coordinates"))
        if len(coordinates) >= 2:
            results.append(("line", coordinates, geometry_properties(element)))

    elif kind == "Polygon":
        outer_ring = None
        inner_rings = []
        for boundary in list(element):
            boundary_name = local_name(boundary.tag)
            if boundary_name not in ("outerBoundaryIs", "innerBoundaryIs"):
                continue
            ring = parse_coordinates(descendant_text(boundary, "coordinates"))
            if len(ring) < 3:
                continue
            if boundary_name == "outerBoundaryIs":
                outer_ring = ring
            else:
                inner_rings.append(ring)

        if outer_ring:
            rings = orient_polygon_rings([outer_ring] + inner_rings)
            results.append(("polygon", rings, geometry_properties(element)))

    elif kind == "Track":
        coordinates = []
        times = []
        for child in list(element):
            child_name = local_name(child.tag)
            if child_name == "coord":
                coordinate = parse_gx_coord(child.text or "")
                if coordinate:
                    coordinates.append(coordinate)
            elif child_name == "when":
                times.append(clean_text(child.text))

        if len(coordinates) >= 2:
            properties = geometry_properties(element)
            if times:
                properties["KML_TRACK_WHEN"] = " | ".join(times)
            results.append(("line", coordinates, properties))

    elif kind in ("MultiGeometry", "MultiTrack"):
        # Her alt geometri bağımsız kayıt olur; tek geometriye birleştirilmez.
        for child in list(element):
            results.extend(extract_geometries(child))

    return results


# -----------------------------------------------------------------------------
# WKT ve temsili koordinat (Excel çıktısı için)
# -----------------------------------------------------------------------------

def _fmt_num(value: float) -> str:
    """WKT için sayıyı gereksiz sıfırlar olmadan yazar."""
    text = repr(float(value))
    return text


def geometry_to_wkt(geometry_type: str, geometry) -> str:
    """
    Geometriyi WKT (Well-Known Text) string olarak döndürür.
        point   -> POINT (x y)
        line    -> LINESTRING (x y, x y, ...)
        polygon -> POLYGON ((dış halka), (iç halka), ...)
    """
    if geometry_type == "point":
        return "POINT ({} {})".format(_fmt_num(geometry[0]), _fmt_num(geometry[1]))

    if geometry_type == "line":
        pairs = ", ".join(
            "{} {}".format(_fmt_num(pt[0]), _fmt_num(pt[1])) for pt in geometry
        )
        return "LINESTRING ({})".format(pairs)

    if geometry_type == "polygon":
        rings = []
        for ring in geometry:
            pairs = ", ".join(
                "{} {}".format(_fmt_num(pt[0]), _fmt_num(pt[1])) for pt in ring
            )
            rings.append("({})".format(pairs))
        return "POLYGON ({})".format(", ".join(rings))

    return ""


def representative_xy(geometry_type: str, geometry):
    """
    Excel'in X ve Y sütunları için temsili koordinat.
        point   -> noktanın kendisi
        line    -> ilk köşe noktası
        polygon -> dış halkanın ilk köşe noktası
    """
    if geometry_type == "point":
        return geometry[0], geometry[1]
    if geometry_type == "line" and geometry:
        return geometry[0][0], geometry[0][1]
    if geometry_type == "polygon" and geometry and geometry[0]:
        return geometry[0][0][0], geometry[0][0][1]
    return "", ""


# -----------------------------------------------------------------------------
# KML öznitelikleri ve katman sınırı
# -----------------------------------------------------------------------------

def collect_schema_map(root):
    schemas = {}
    for element in root.iter():
        if local_name(element.tag) != "Schema":
            continue
        schema_id = element.attrib.get("id") or element.attrib.get("name") or ""
        if not schema_id:
            continue
        fields = {}
        for child in list(element):
            if local_name(child.tag) != "SimpleField":
                continue
            name = clean_text(child.attrib.get("name"))
            display_name = direct_child_text(child, "displayName")
            if name:
                fields[name] = display_name or name
        schemas[schema_id] = fields
    return schemas


def extract_extended_data(placemark, schema_map) -> OrderedDict:
    attributes = OrderedDict()
    for element in placemark.iter():
        element_name = local_name(element.tag)

        if element_name == "Data":
            raw_name = clean_text(element.attrib.get("name"))
            display_name = direct_child_text(element, "displayName")
            value = direct_child_text(element, "value")
            key = display_name or raw_name
            if key:
                attributes[key] = value

        elif element_name == "SchemaData":
            schema_url = clean_text(element.attrib.get("schemaUrl")).lstrip("#")
            field_names = schema_map.get(schema_url, {})
            for child in list(element):
                if local_name(child.tag) != "SimpleData":
                    continue
                raw_name = clean_text(child.attrib.get("name"))
                key = field_names.get(raw_name, raw_name)
                if key:
                    attributes[key] = clean_text(child.text)

        elif element_name == "SimpleArrayData":
            raw_name = clean_text(element.attrib.get("name"))
            values = [
                clean_text(child.text)
                for child in list(element)
                if local_name(child.tag) == "value"
            ]
            if raw_name:
                attributes[raw_name] = " | ".join(value for value in values if value)

    return attributes


def extract_common_attributes(placemark, layer_path: str, source_kml: str, schema_map):
    attributes = OrderedDict()

    for name in (
        "name", "description", "address", "phoneNumber", "visibility",
        "open", "styleUrl", "Snippet",
    ):
        value = direct_child_text(placemark, name)
        if value:
            attributes["KML_{}".format(name)] = value

    attributes["KML_LAYER"] = layer_path
    attributes["KML_SOURCE"] = source_kml

    timestamp = ""
    time_begin = ""
    time_end = ""
    for element in placemark.iter():
        element_name = local_name(element.tag)
        if element_name == "TimeStamp":
            timestamp = direct_child_text(element, "when")
        elif element_name == "TimeSpan":
            time_begin = direct_child_text(element, "begin")
            time_end = direct_child_text(element, "end")

    if timestamp:
        attributes["KML_TIME"] = timestamp
    if time_begin:
        attributes["KML_BEGIN"] = time_begin
    if time_end:
        attributes["KML_END"] = time_end

    description = attributes.get("KML_description", "")
    for key, value in parse_description_table(description).items():
        if key not in attributes:
            attributes[key] = value
        else:
            attributes["Popup_{}".format(key)] = value

    for key, value in extract_extended_data(placemark, schema_map).items():
        if key not in attributes:
            attributes[key] = value
        elif value != attributes[key]:
            attributes["Extended_{}".format(key)] = value

    return attributes


def find_geometry_children(placemark):
    geometry_names = {"Point", "LineString", "Polygon", "MultiGeometry", "Track", "MultiTrack"}
    return [child for child in list(placemark) if local_name(child.tag) in geometry_names]


def current_layer_path(source_kml: str, document_name: str, folders: list[str]) -> str:
    """Katman kimliği: KML üyesi + document + tam folder yolu."""
    parts = []
    if document_name:
        parts.append(document_name)
    parts.extend([item for item in folders if item])
    if not parts:
        parts.append(Path(source_kml).stem or "katman")
    return "/".join(parts)


def walk_kml(
    node,
    folders,
    source_kml,
    document_name,
    schema_map,
    grouped_records,
    placemark_counter,
):
    for child in list(node):
        element_name = local_name(child.tag)

        if element_name == "Folder":
            folder_name = direct_child_text(child, "name") or "adsiz_klasor"
            walk_kml(
                child,
                folders + [folder_name],
                source_kml,
                document_name,
                schema_map,
                grouped_records,
                placemark_counter,
            )

        elif element_name == "Document":
            child_document_name = direct_child_text(child, "name") or document_name
            walk_kml(
                child,
                folders,
                source_kml,
                child_document_name,
                schema_map,
                grouped_records,
                placemark_counter,
            )

        elif element_name == "Placemark":
            placemark_counter[0] += 1
            placemark_id = clean_text(child.attrib.get("id")) or "PM_{:07d}".format(placemark_counter[0])
            layer_path = current_layer_path(source_kml, document_name, folders)
            attributes = extract_common_attributes(child, layer_path, source_kml, schema_map)

            geometry_items = []
            for geometry_element in find_geometry_children(child):
                geometry_items.extend(extract_geometries(geometry_element))

            # Aynı Placemark içindeki geometri örnekleri ayrı kayıt olarak tutulur.
            geometry_type_counts = defaultdict(int)
            for geometry_type, geometry, geometry_attributes in geometry_items:
                geometry_type_counts[geometry_type] += 1
                merged_attributes = OrderedDict(attributes)
                merged_attributes.update(geometry_attributes)

                group_key = (source_kml, layer_path, geometry_type)
                grouped_records[group_key].append(
                    {
                        "src_id": placemark_id,
                        "part_no": geometry_type_counts[geometry_type],
                        "attrs": merged_attributes,
                        "geometry": geometry,
                    }
                )

        elif list(child):
            walk_kml(
                child,
                folders,
                source_kml,
                document_name,
                schema_map,
                grouped_records,
                placemark_counter,
            )


def load_kmz(kmz_path: Path):
    grouped_records = defaultdict(list)
    placemark_counter = [0]

    with zipfile.ZipFile(kmz_path, "r") as archive:
        kml_members = [name for name in archive.namelist() if name.lower().endswith(".kml")]
        if not kml_members:
            raise ValueError("KMZ içerisinde KML dosyası bulunamadı.")

        kml_members.sort(key=lambda name: (Path(name).name.lower() != "doc.kml", name.lower()))

        for member in kml_members:
            raw_xml = archive.read(member)
            root = ET.fromstring(raw_xml)
            schema_map = collect_schema_map(root)
            root_document_name = direct_child_text(root, "name")
            walk_kml(
                root,
                folders=[],
                source_kml=member,
                document_name=root_document_name,
                schema_map=schema_map,
                grouped_records=grouped_records,
                placemark_counter=placemark_counter,
            )

    return grouped_records


def load_kml(kml_path: Path):
    """Tek bir KML dosyasını, KMZ içindeki KML ile aynı katman mantığında oku."""
    grouped_records = defaultdict(list)
    placemark_counter = [0]
    root = ET.parse(str(kml_path)).getroot()
    schema_map = collect_schema_map(root)
    root_document_name = direct_child_text(root, "name")
    walk_kml(
        root,
        folders=[],
        source_kml=kml_path.name,
        document_name=root_document_name,
        schema_map=schema_map,
        grouped_records=grouped_records,
        placemark_counter=placemark_counter,
    )
    return grouped_records


# -----------------------------------------------------------------------------
# DBF alan eşlemesi ve yan dosyalar
# -----------------------------------------------------------------------------

def ordered_attribute_keys(records):
    keys = OrderedDict()
    for record in records:
        for key in record["attrs"].keys():
            keys.setdefault(key, None)
    return list(keys.keys())


def build_dbf_mapping(records, attribute_keys):
    mapping = []
    csv_only = []
    used_dbf_names = {"src_id", "part_no"}
    remaining_fields = DBF_MAX_FIELDS - 2

    for key in attribute_keys:
        maximum_length = max(
            (len(clean_text(record["attrs"].get(key, ""))) for record in records),
            default=0,
        )
        chunk_count = max(1, int(math.ceil(maximum_length / float(DBF_TEXT_LIMIT))))

        if chunk_count > remaining_fields:
            csv_only.append(key)
            continue

        base_seed = readable_dbf_seed(key)
        for chunk_index in range(chunk_count):
            chunk_seed = base_seed if chunk_count == 1 else "{}_{}".format(base_seed, chunk_index + 1)
            dbf_name = unique_dbf_name(chunk_seed, used_dbf_names)
            mapping.append(
                {
                    "dbf_name": dbf_name,
                    "original_name": key,
                    "alias": key if chunk_count == 1 else "{} [{}/{}]".format(key, chunk_index + 1, chunk_count),
                    "chunk_index": chunk_index + 1,
                    "chunk_count": chunk_count,
                    "start": chunk_index * DBF_TEXT_LIMIT,
                    "end": (chunk_index + 1) * DBF_TEXT_LIMIT,
                }
            )
            remaining_fields -= 1

    return mapping, csv_only


def write_qml_aliases(output_base: Path, mapping):
    aliases = [
        ("SRC_ID", "Kaynak Placemark kimliği"),
        ("PART_NO", "Placemark içindeki geometri sıra numarası"),
    ]
    aliases.extend((item["dbf_name"], item["alias"]) for item in mapping)

    lines = [
        '<!DOCTYPE qgis PUBLIC "http://mrcc.com/qgis.dtd" "SYSTEM">',
        '<qgis version="3.0.0" styleCategories="Fields">',
        "  <aliases>",
    ]
    for index, (field_name, alias) in enumerate(aliases):
        lines.append(
            '    <alias field="{}" index="{}" name="{}"/>'.format(
                xml_escape(field_name, {'"': '&quot;'}),
                index,
                xml_escape(alias, {'"': '&quot;'}),
            )
        )
    lines.extend(["  </aliases>", "</qgis>"])
    output_base.with_suffix(".qml").write_text("\n".join(lines), encoding="utf-8")


def write_sidecars(output_base, records, attribute_keys, mapping, csv_only):
    attributes_csv = output_base.with_name(output_base.name + "_attributes.csv")
    with attributes_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["SRC_ID", "PART_NO"] + attribute_keys,
            extrasaction="ignore",
        )
        writer.writeheader()
        for record in records:
            row = {"SRC_ID": record["src_id"], "PART_NO": record["part_no"]}
            row.update({key: clean_text(record["attrs"].get(key, "")) for key in attribute_keys})
            writer.writerow(row)

    mapping_csv = output_base.with_name(output_base.name + "_field_map.csv")
    with mapping_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        fieldnames = [
            "DBF_FIELD", "ORIGINAL_FIELD", "DISPLAY_ALIAS", "CHUNK_NO",
            "CHUNK_COUNT", "CHAR_START", "CHAR_END", "STORAGE",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for item in mapping:
            writer.writerow(
                {
                    "DBF_FIELD": item["dbf_name"],
                    "ORIGINAL_FIELD": item["original_name"],
                    "DISPLAY_ALIAS": item["alias"],
                    "CHUNK_NO": item["chunk_index"],
                    "CHUNK_COUNT": item["chunk_count"],
                    "CHAR_START": item["start"] + 1,
                    "CHAR_END": item["end"],
                    "STORAGE": "DBF_AND_CSV",
                }
            )
        for key in csv_only:
            writer.writerow(
                {
                    "DBF_FIELD": "",
                    "ORIGINAL_FIELD": key,
                    "DISPLAY_ALIAS": key,
                    "CHUNK_NO": "",
                    "CHUNK_COUNT": "",
                    "CHAR_START": "",
                    "CHAR_END": "",
                    "STORAGE": "CSV_ONLY",
                }
            )

    jsonl_path = output_base.with_name(output_base.name + "_attributes.jsonl")
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for record in records:
            payload = OrderedDict(
                [
                    ("SRC_ID", record["src_id"]),
                    ("PART_NO", record["part_no"]),
                    ("attributes", record["attrs"]),
                ]
            )
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")


def _truncate_utf8(text: str, max_bytes: int) -> str:
    """DBF alanı bayt sınırını aşmasın diye metni UTF-8 bayt bazında kırpar.
    (Tam/kesilmemiş değerler Excel, *_attributes.csv ve *.jsonl içinde durur.)"""
    encoded = text.encode("utf-8")
    if len(encoded) <= max_bytes:
        return text
    return encoded[:max_bytes].decode("utf-8", "ignore")


def write_shapefile(output_base: Path, geometry_type: str, records):
    if not records:
        return

    shape_types = {
        "point": shapefile.POINT,
        "line": shapefile.POLYLINE,
        "polygon": shapefile.POLYGON,
    }

    attribute_keys = ordered_attribute_keys(records)
    mapping, csv_only = build_dbf_mapping(records, attribute_keys)

    writer = shapefile.Writer(
        str(output_base),
        shapeType=shape_types[geometry_type],
        encoding="utf-8",
        encodingErrors="strict",
    )
    writer.autoBalance = 1
    writer.field("SRC_ID", "C", size=80)
    writer.field("PART_NO", "N", size=10, decimal=0)

    for item in mapping:
        writer.field(item["dbf_name"], "C", size=DBF_TEXT_LIMIT)

    for record in records:
        geometry = record["geometry"]

        # Her kayıt için tam olarak bir Shape yazılır; geometri birleştirme yoktur.
        if geometry_type == "point":
            writer.point(geometry[0], geometry[1])
        elif geometry_type == "line":
            writer.line([[[point[0], point[1]] for point in geometry]])
        elif geometry_type == "polygon":
            writer.poly(
                [
                    [[point[0], point[1]] for point in ring]
                    for ring in geometry
                ]
            )

        values = [record["src_id"][:80], int(record["part_no"])]
        for item in mapping:
            text = clean_text(record["attrs"].get(item["original_name"], ""))
            values.append(_truncate_utf8(text[item["start"] : item["end"]], DBF_TEXT_LIMIT))
        writer.record(*values)

    writer.close()
    output_base.with_suffix(".cpg").write_text("UTF-8", encoding="ascii")
    output_base.with_suffix(".prj").write_text(WGS84_PRJ, encoding="ascii")
    write_qml_aliases(output_base, mapping)
    write_sidecars(output_base, records, attribute_keys, mapping, csv_only)


# -----------------------------------------------------------------------------
# Excel çıktısı (X, Y ve WKT geometri dahil)
# -----------------------------------------------------------------------------

def write_excel_workbook(excel_path: Path, sorted_groups):
    """
    Tüm katmanları tek bir Excel çalışma kitabına yazar.

    - "Tümü" sayfası: tüm kayıtlar tek tabloda (LAYER + GEOMETRY_TYPE sütunlu).
    - Her sütun: SRC_ID, PART_NO, LAYER, GEOMETRY_TYPE, X, Y, WKT + öznitelikler.
    - X ve Y temsili koordinat; WKT tam geometri (poligon ise POLYGON(...) vb.).
    - Öznitelikler tüm katmanların birleşiminden gelen sütunlarda toplanır.
    """
    if not HAVE_OPENPYXL:
        print("UYARI: openpyxl kurulu değil, Excel çıktısı atlandı. "
              "Kurmak için: python -m pip install openpyxl")
        return None

    # Tüm kayıtlar ve öznitelik anahtarlarının birleşimi.
    all_records = []
    attribute_keys = OrderedDict()
    for (source_kml, layer_path, geometry_type), records in sorted_groups:
        for record in records:
            for key in record["attrs"].keys():
                attribute_keys.setdefault(key, None)
            all_records.append((layer_path, geometry_type, record))

    attribute_keys = list(attribute_keys.keys())

    base_headers = ["SRC_ID", "PART_NO", "LAYER", "GEOMETRY_TYPE", "X", "Y", "WKT"]
    headers = base_headers + attribute_keys

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Tümü"
    sheet.append(headers)

    for layer_path, geometry_type, record in all_records:
        geometry = record["geometry"]
        x, y = representative_xy(geometry_type, geometry)
        wkt = geometry_to_wkt(geometry_type, geometry)

        row = [
            record["src_id"],
            record["part_no"],
            layer_path,
            geometry_type,
            x,
            y,
            wkt,
        ]
        for key in attribute_keys:
            row.append(clean_text(record["attrs"].get(key, "")))
        sheet.append(row)

    # Başlık satırını dondur ve makul sütun genişlikleri ver.
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        if header == "WKT":
            width = 60
        elif header in ("LAYER", "SRC_ID"):
            width = 28
        else:
            width = max(10, min(40, len(str(header)) + 4))
        sheet.column_dimensions[letter].width = width

    workbook.save(str(excel_path))
    return excel_path


# -----------------------------------------------------------------------------
# Ana dönüşüm
# -----------------------------------------------------------------------------

def convert_kmz_to_layer_shapefiles(kmz_file, output_folder, write_excel=True):
    kmz_path = Path(kmz_file).expanduser().resolve()
    output_path = Path(output_folder).expanduser().resolve()

    if not kmz_path.exists():
        raise FileNotFoundError("KMZ bulunamadı: {}".format(kmz_path))
    if kmz_path.suffix.casefold() not in {".kml", ".kmz"}:
        raise ValueError("Girdi dosyası .kml veya .kmz olmalıdır: {}".format(kmz_path))

    output_path.mkdir(parents=True, exist_ok=True)
    grouped_records = load_kml(kmz_path) if kmz_path.suffix.casefold() == ".kml" else load_kmz(kmz_path)

    if not grouped_records:
        raise ValueError("KMZ içerisinde desteklenen geometri bulunamadı.")

    prefix = safe_file_part(kmz_path.stem, fallback="kmz")
    used_output_names = set()
    manifest_rows = []

    sorted_groups = sorted(
        grouped_records.items(),
        key=lambda item: (item[0][0].casefold(), item[0][1].casefold(), item[0][2]),
    )

    for (source_kml, layer_path, geometry_type), records in sorted_groups:
        layer_leaf = layer_path.split("/")[-1] if layer_path else Path(source_kml).stem
        seed = "{}_{}_{}".format(prefix, layer_leaf, GEOMETRY_SUFFIX[geometry_type])
        output_name = make_unique_output_name(seed, used_output_names)
        output_base = output_path / output_name

        write_shapefile(output_base, geometry_type, records)

        manifest_rows.append(
            {
                "OUTPUT_SHP": output_base.name + ".shp",
                "SOURCE_KML": source_kml,
                "SOURCE_LAYER_PATH": layer_path,
                "GEOMETRY_TYPE": geometry_type,
                "FEATURE_COUNT": len(records),
                "DISSOLVE": "NO",
                "MERGE_WITH_OTHER_LAYER": "NO",
            }
        )

    excel_path = None
    if write_excel:
        excel_path = write_excel_workbook(
            output_path / "{}_veriler.xlsx".format(prefix),
            sorted_groups,
        )

    manifest_path = output_path / "{}_layer_manifest.csv".format(prefix)
    with manifest_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "OUTPUT_SHP", "SOURCE_KML", "SOURCE_LAYER_PATH", "GEOMETRY_TYPE",
                "FEATURE_COUNT", "DISSOLVE", "MERGE_WITH_OTHER_LAYER",
            ],
        )
        writer.writeheader()
        writer.writerows(manifest_rows)

    print("Dönüştürme tamamlandı. Katmanlar birleştirilmedi:")
    for row in manifest_rows:
        print(
            " - {} | {} | {} obje".format(
                row["OUTPUT_SHP"],
                row["SOURCE_LAYER_PATH"],
                row["FEATURE_COUNT"],
            )
        )
    print("\nKatman eşlemesi: {}".format(manifest_path))
    print("QGIS alan alias dosyaları: aynı adlı *.qml")
    print("Tam öznitelikler: *_attributes.csv ve *_attributes.jsonl")
    if excel_path:
        print("Excel çıktısı (X, Y, WKT dahil): {}".format(excel_path))


def _silence_dbf_truncation_warnings():
    """DBF 254 baytlık alan sınırı nedeniyle uzun HTML açıklama alanları kırpılır.
    Bu kasıtlıdır; tam veri Excel, *_attributes.csv ve *.jsonl içinde saklanır.
    pyshp'in PossibleDataLoss uyarısı bu yüzden bastırılır."""
    try:
        from shapefile import PossibleDataLoss
        warnings.filterwarnings("ignore", category=PossibleDataLoss)
    except Exception:
        pass


def _force_utf8_stdout():
    """Windows konsolu (cp1252) Türkçe karakterlerde çökmesin diye UTF-8'e geçir."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


def main():
    _force_utf8_stdout()
    _silence_dbf_truncation_warnings()
    parser = argparse.ArgumentParser(
        description=(
            "KML/KMZ dosyasını ArcPy/GDAL olmadan, her KML katmanını ayrı tutarak "
            "Shapefile formatına dönüştürür."
        )
    )
    parser.add_argument("kmz", help="Girdi KML veya KMZ dosyasının yolu")
    parser.add_argument("output_folder", help="Shapefile çıktı klasörü")
    parser.add_argument(
        "--no-excel",
        action="store_true",
        help="Excel (.xlsx) çıktısı üretme; yalnızca Shapefile ve yan dosyalar.",
    )
    args = parser.parse_args()
    convert_kmz_to_layer_shapefiles(
        args.kmz, args.output_folder, write_excel=not args.no_excel
    )


if __name__ == "__main__":
    main()
